from flask import Blueprint, render_template, redirect, url_for, flash, request, abort, jsonify
from flask_login import login_required, current_user
from sqlalchemy import or_
from app import db
from app.models import Document, DocumentItem, DocumentComment, DocumentApproval, DOC_TYPES, DOC_STATUSES, User, Notification, DocumentAttachment
from app.decorators import requires_permission
from app.audit import log_action
from app.storage import save_upload, send_attachment, allowed_file, MAX_FILE_SIZE_MB
from datetime import datetime

documents = Blueprint('documents', __name__, url_prefix='/documents',
                      template_folder='../../app/templates/documents')


def _to_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _unpack_extras(justification):
    """Parse the packed 'Label: Value' lines back into form-field names."""
    label_to_field = {
        'Заказано': 'ordered_by',
        'Плательщик НДС': 'vat_payer',
        'Валюта': 'currency',
        'Дата оплаты': 'payment_date',
        'Основание': 'basis',
        'Контрагент': 'counterparty',
        'Условия оплаты': 'payment_terms',
        'Условия оказания услуг': 'service_terms',
        'Срок оказания услуг': 'service_date',
        'Статья бюджета': 'budget_line',
        'Требование на приобретение материалов': 'material_request',
    }
    out = {}
    for line in (justification or '').split('\n'):
        if ': ' in line:
            label, val = line.split(': ', 1)
            field = label_to_field.get(label.strip())
            if field:
                out[field] = val.strip()
    return out


def _assigned_doc_ids_subquery(user):
    """Document IDs where this user is (or was) an assigned approver at any step —
    regardless of their role's blanket permissions. Being routed to a document
    as a reviewer/signatory is its own authorization to see and act on it."""
    return db.session.query(DocumentApproval.document_id).filter_by(approver_id=user.id)


def _is_assigned_approver(doc, user):
    return DocumentApproval.query.filter_by(document_id=doc.id, approver_id=user.id).first() is not None


PROCUREMENT_DOC_TYPES = ('purchase_req', 'po_services')


def _can_view_doc(user, doc):
    """View access: broad permission, own doc, routed approver, or procurement
    for procurement document types."""
    if user.can_access('documents') or user.can_access('documents_read'):
        return True
    if doc.author_id == user.id or _is_assigned_approver(doc, user):
        return True
    return (doc.doc_type in PROCUREMENT_DOC_TYPES
            and user.can_access('documents_procurement'))


def _visible_docs_query(user):
    """Documents a user may see: everything (broad 'documents' permission),
    or their own + anything they're routed to approve/review + procurement types
    for procurement staff."""
    query = Document.query
    if user.can_access('documents') or user.can_access('documents_read'):
        return query
    conditions = [
        Document.author_id == user.id,
        Document.id.in_(_assigned_doc_ids_subquery(user)),
    ]
    if user.can_access('documents_procurement'):
        conditions.append(Document.doc_type.in_(PROCUREMENT_DOC_TYPES))
    return query.filter(or_(*conditions))


def _panel_counts(doc_type=None):
    """Status counters for the left panel via a single aggregate query —
    stays fast no matter how many documents accumulate over the years."""
    from sqlalchemy import func
    visible = _visible_docs_query(current_user).subquery()
    query = db.session.query(visible.c.status, func.count()).group_by(visible.c.status)
    if doc_type:
        query = query.filter(visible.c.doc_type == doc_type)
    counts = dict(query.all())
    counts['_total'] = sum(counts.values())
    return counts


def _build_route(doc, action, signatory_id):
    """Create DocumentApproval records. Shared by all forms.

    Order: reviewer stages first (steps 0..k-1, in form order), the signatory
    signs LAST (step k) — Создание → Согласование → Подпись.

    New form format: each stage emits stage_no[] (aligned with stage_type[])
    and its reviewers as stage_reviewer_<no>[]. Legacy flat stage_reviewer[]
    is treated as a single parallel stage.
    """
    if action != 'submit' or not signatory_id:
        return
    stage_nos = request.form.getlist('stage_no[]')
    stages = []
    if stage_nos:
        for no in stage_nos:
            ids = [int(r) for r in request.form.getlist(f'stage_reviewer_{no}[]')
                   if r.strip().isdigit()]
            if ids:
                stages.append(ids)
    else:
        ids = [int(r) for r in request.form.getlist('stage_reviewer[]')
               if r.strip().isdigit()]
        if ids:
            stages.append(ids)

    step = 0
    for ids in stages:
        for rid in ids:
            db.session.add(DocumentApproval(document_id=doc.id, approver_id=rid,
                                            step=step, status='pending'))
        step += 1
    # the signatory always signs last
    db.session.add(DocumentApproval(document_id=doc.id, approver_id=signatory_id,
                                    step=step, status='pending'))


def _save_form_attachments(doc):
    """Save files posted in the 'attachments' field of a create/edit form.
    Mirrors the validation in upload_attachment()."""
    from werkzeug.utils import secure_filename
    for f in request.files.getlist('attachments'):
        if not f or not f.filename:
            continue
        if not allowed_file(f.filename):
            flash(f'Файл «{f.filename}» не загружен: недопустимый тип.', 'warning')
            continue
        stored_filename, backend, size_bytes = save_upload(f)
        if size_bytes > MAX_FILE_SIZE_MB * 1024 * 1024:
            flash(f'Файл «{f.filename}» не загружен: больше {MAX_FILE_SIZE_MB} МБ.', 'warning')
            continue
        db.session.add(DocumentAttachment(
            document_id=doc.id,
            original_filename=secure_filename(f.filename),
            stored_filename=stored_filename,
            storage_backend=backend,
            content_type=f.content_type,
            size_bytes=size_bytes,
            uploaded_by=current_user.id,
        ))


def _notify_approvers(doc):
    all_approvals = DocumentApproval.query.filter_by(document_id=doc.id).all()
    notified = set()
    for appr in all_approvals:
        if appr.approver_id not in notified:
            db.session.add(Notification(
                user_id=appr.approver_id,
                title=f'Документ на согласование: {doc.doc_number}',
                body=doc.title[:100],
                link=f'/documents/{doc.id}',
                is_read=False,
            ))
            notified.add(appr.approver_id)


def _notify_current_approvers(doc, title):
    """Notify whoever is pending at the document's current step."""
    pend = DocumentApproval.query.filter_by(
        document_id=doc.id, step=doc.current_step, status='pending').all()
    for appr in pend:
        db.session.add(Notification(
            user_id=appr.approver_id,
            title=title,
            body=doc.title[:100],
            link=f'/documents/{doc.id}',
            is_read=False,
        ))


@documents.route('/')
@login_required
def index():
    status = request.args.get('status')
    year = request.args.get('year', type=int)
    doc_type = request.args.get('doc_type')
    page = request.args.get('page', 1, type=int)

    panel_counts = _panel_counts(doc_type)

    query = _visible_docs_query(current_user)
    if status:
        query = query.filter_by(status=status)
    if doc_type:
        query = query.filter_by(doc_type=doc_type)
    if year:
        from sqlalchemy import extract
        query = query.filter(extract('year', Document.created_at) == year)
    query = query.order_by(Document.created_at.desc())

    pagination = query.paginate(page=page, per_page=25, error_out=False)
    docs = pagination.items

    from sqlalchemy import extract, distinct
    years = sorted({y for (y,) in db.session.query(
        distinct(extract('year', Document.created_at))).all() if y}, reverse=True)
    return render_template('documents/index.html', docs=docs, panel_counts=panel_counts,
                           years=years, sel_year=year, sel_type=doc_type,
                           pagination=pagination,
                           doc_types=DOC_TYPES, statuses=DOC_STATUSES)

@documents.route('/new')
@login_required
def new():
    return render_template('documents/new.html', doc_types=DOC_TYPES)


@documents.route('/new/purchase-requisition', methods=['GET', 'POST'])
@login_required
def new_purchase_req():
    """Old wizard — funnels into the single full form (kept for old links/tests)."""
    if request.method == 'POST':
        return submit_trebovanie_new()
    return redirect(url_for('documents.new_trebovanie_new',
                            from_defect=request.args.get('from_defect')))


@documents.route('/<int:doc_id>')
@login_required
def view(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if not _can_view_doc(current_user, doc):
        abort(403)
    items     = doc.items.all()
    comments  = doc.comments.order_by('created_at').all()
    approvals = doc.approvals.order_by(DocumentApproval.step).all()
    panel_counts = _panel_counts()
    return render_template('documents/view.html', doc=doc, items=items,
                           comments=comments, approvals=approvals, panel_counts=panel_counts)


@documents.route('/<int:doc_id>/attachments/upload', methods=['POST'])
@login_required
def upload_attachment(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if not current_user.can_access('documents') and doc.author_id != current_user.id:
        abort(403)

    file = request.files.get('file')
    if not file or not file.filename:
        flash('Please choose a file.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc.id))

    if not allowed_file(file.filename):
        flash('That file type is not allowed.', 'danger')
        return redirect(url_for('documents.view', doc_id=doc.id))

    stored_filename, backend, size_bytes = save_upload(file)
    if size_bytes > MAX_FILE_SIZE_MB * 1024 * 1024:
        flash(f'File is larger than {MAX_FILE_SIZE_MB} MB.', 'danger')
        return redirect(url_for('documents.view', doc_id=doc.id))

    from werkzeug.utils import secure_filename
    attachment = DocumentAttachment(
        document_id=doc.id,
        original_filename=secure_filename(file.filename),
        stored_filename=stored_filename,
        storage_backend=backend,
        content_type=file.content_type,
        size_bytes=size_bytes,
        uploaded_by=current_user.id,
    )
    db.session.add(attachment)
    db.session.commit()
    log_action('attachment_uploaded', 'document', doc.id, details=attachment.original_filename)
    flash('File attached.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/<int:doc_id>/attachments/<int:attachment_id>/download')
@login_required
def download_attachment(doc_id, attachment_id):
    doc = Document.query.get_or_404(doc_id)
    if not _can_view_doc(current_user, doc):
        abort(403)
    attachment = DocumentAttachment.query.filter_by(id=attachment_id, document_id=doc.id).first_or_404()
    return send_attachment(attachment)


@documents.route('/<int:doc_id>/attachments/<int:attachment_id>/delete', methods=['POST'])
@login_required
def delete_attachment(doc_id, attachment_id):
    doc = Document.query.get_or_404(doc_id)
    if not current_user.can_access('documents') and doc.author_id != current_user.id:
        abort(403)
    attachment = DocumentAttachment.query.filter_by(id=attachment_id, document_id=doc.id).first_or_404()
    db.session.delete(attachment)
    db.session.commit()
    flash('Attachment removed.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/<int:doc_id>/approve', methods=['POST'])
@login_required
def approve(doc_id):
    # Lock the document row for the duration of this decision so two approvers
    # acting on the same parallel stage at the same instant can't both see
    # "I'm the last pending one" and double-advance current_step. Row locking
    # is a Postgres-only concern here (production); SQLite (used in tests)
    # doesn't support FOR UPDATE, so skip it there rather than fail loudly.
    _lockable = db.session.get_bind().dialect.name == 'postgresql'
    doc_query = Document.query.filter_by(id=doc_id)
    doc = (doc_query.with_for_update().first() if _lockable else doc_query.first())
    if doc is None:
        abort(404)
    action = request.form.get('action')
    comment_text = request.form.get('comment', '').strip()

    if action in ['approve', 'reject', 'return']:
        # No blanket role check here on purpose: being the assigned approver for
        # this document's current step (checked just below) is the real
        # authorization, regardless of the approver's role. A finance/HR/mechanic
        # user who was specifically routed to review this document must be able
        # to act on it even though their role has no broad 'documents' permission.
        if doc.status != 'pending':
            flash('Этот документ уже обработан и больше не ожидает согласования.', 'warning')
            return redirect(url_for('documents.view', doc_id=doc_id))

        approval_query = DocumentApproval.query.filter_by(
            document_id=doc_id,
            approver_id=current_user.id,
            step=doc.current_step,
            status='pending'
        )
        approval = (approval_query.with_for_update().first() if _lockable else approval_query.first())

        if not approval and current_user.role == 'it_admin':
            fallback_query = DocumentApproval.query.filter_by(
                document_id=doc_id,
                step=doc.current_step,
                status='pending'
            )
            approval = (fallback_query.with_for_update().first() if _lockable else fallback_query.first())

        if not approval:
            flash('Вы не можете выполнить это действие для данного документа.', 'warning')
            return redirect(url_for('documents.view', doc_id=doc_id))

        if action == 'return' and not comment_text:
            flash('Укажите причину возврата документа.', 'warning')
            return redirect(url_for('documents.view', doc_id=doc_id))
        if action == 'reject' and not comment_text:
            flash('Укажите причину отклонения документа.', 'warning')
            return redirect(url_for('documents.view', doc_id=doc_id))

        if action == 'return':
            doc.status = 'returned'
            db.session.add(DocumentComment(
                document_id=doc.id,
                author_id=current_user.id,
                text=f'Документ возвращён на доработку: {comment_text}',
                is_system=False,
            ))
            db.session.add(Notification(
                user_id=doc.author_id,
                title=f'Документ {doc.doc_number}: возвращён на доработку',
                body=comment_text[:100],
                link=f'/documents/{doc.id}',
                is_read=False,
            ))
            log_action('document_returned', 'document', doc.id, details=doc.doc_number)
            db.session.commit()
            flash('Документ возвращён автору на доработку.', 'success')
            return redirect(url_for('documents.view', doc_id=doc_id))

        approval.status = 'approved' if action == 'approve' else 'rejected'
        approval.decided_at = datetime.utcnow()
        if comment_text:
            approval.comment = comment_text

        if action == 'reject':
            doc.status = 'rejected'
            doc.current_step = -1
            # Close out any other still-pending approvals (other reviewers in this
            # stage, or later stages/signatory) so the trail doesn't show phantom
            # "awaiting approval" items on a document that's already dead.
            DocumentApproval.query.filter_by(document_id=doc_id, status='pending')\
                .update({'status': 'skipped'}, synchronize_session=False)
        else:
            db.session.flush()   # decision above must be visible to the count
            pending_at_step = DocumentApproval.query.filter_by(
                document_id=doc_id,
                step=doc.current_step,
                status='pending'
            ).count()

            if pending_at_step == 0:
                next_step_approvals = DocumentApproval.query.filter_by(
                    document_id=doc_id,
                    step=doc.current_step + 1
                ).count()

                if next_step_approvals > 0:
                    doc.current_step += 1
                    _notify_current_approvers(doc, f'Документ на согласование: {doc.doc_number}')
                else:
                    doc.status = 'approved'
                    if doc.doc_type == 'purchase_req':
                        for proc in User.query.filter_by(role='procurement', is_active=True).all():
                            db.session.add(Notification(
                                user_id=proc.id,
                                title=f'Требование подписано: {doc.doc_number}',
                                body=(doc.title or '')[:100],
                                link=f'/documents/{doc.id}', is_read=False))

        text = comment_text or f'Документ {"согласован" if action == "approve" else "отклонён"} пользователем {current_user.full_name}.'
        db.session.add(DocumentComment(
            document_id=doc.id,
            author_id=current_user.id,
            text=text,
            is_system=False,
        ))

        db.session.add(Notification(
            user_id=doc.author_id,
            title=f'Документ {doc.doc_number}: {"согласован" if action == "approve" else "отклонён"}',
            body=doc.title[:100],
            link=f'/documents/{doc.id}',
            is_read=False,
        ))

        log_action(f'document_{"approved" if action == "approve" else "rejected"}',
                   'document', doc.id, details=doc.doc_number)
        db.session.commit()
        flash(f'Документ {"согласован" if action == "approve" else "отклонён"}.', 'success')

    elif action == 'comment' and comment_text:
        if (doc.author_id != current_user.id
                and not current_user.can_access('documents')
                and not _is_assigned_approver(doc, current_user)):
            abort(403)
        db.session.add(DocumentComment(
            document_id=doc.id,
            author_id=current_user.id,
            text=comment_text,
            is_system=False,
        ))
        db.session.commit()

    return redirect(url_for('documents.view', doc_id=doc_id))


@documents.route('/<int:doc_id>/resubmit', methods=['POST'])
@login_required
def resubmit(doc_id):
    """Author re-sends a returned document back to the supervisor."""
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.status != 'returned':
        flash('Этот документ нельзя повторно отправить.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    doc.status = 'pending'
    db.session.add(DocumentComment(
        document_id=doc.id,
        author_id=current_user.id,
        text=f'Документ повторно отправлен на согласование пользователем {current_user.full_name}.',
        is_system=True,
    ))
    _notify_current_approvers(doc, f'Документ повторно на согласовании: {doc.doc_number}')
    log_action('document_resubmitted', 'document', doc.id, details=doc.doc_number)
    db.session.commit()
    flash('Документ повторно отправлен на согласование.', 'success')
    return redirect(url_for('documents.view', doc_id=doc_id))


@documents.route('/defect-act/new', methods=['GET', 'POST'])
@login_required
def new_defect_act():
    from app.models import Equipment
    equipment = Equipment.query.order_by(Equipment.unit_id).all()
    now = datetime.now().strftime('%d.%m.%Y %H:%M')
    return render_template('documents/defect.html', equipment=equipment, now=now)


@documents.route('/defect-act/submit', methods=['POST'])
@login_required
def submit_defect_act():
    action = request.form.get('action', 'draft')

    # Auto-assign the approving signatory: department head, else director, else IT admin
    signatory = (User.query.filter_by(role='dept_head', is_active=True).first()
                 or User.query.filter_by(role='director', is_active=True).first()
                 or User.query.filter_by(role='it_admin', is_active=True).first())

    if action == 'submit' and not signatory:
        action = 'draft'
        flash('Документ сохранён как черновик: не найден ответственный для согласования (нет активного руководителя подразделения, директора или IT-админа). Обратитесь к администратору.', 'warning')

    equipment_id = request.form.get('equipment_id', type=int)
    doc = Document(
        doc_type      = 'defect_act',
        equipment_id  = equipment_id,
        title         = request.form.get('description', 'Дефектный акт')[:100],
        department    = request.form.get('department', current_user.department),
        urgency       = request.form.get('urgency', 'critical'),
        purpose       = request.form.get('description'),
        justification = request.form.get('cause'),
        author_id     = current_user.id,
        status        = 'pending' if action == 'submit' else 'draft',
        current_step  = 0,
    )

    db.session.add(doc)
    db.session.flush()
    doc.generate_number()
    doc.assign_event_code()
    log_action('document_created', 'document', doc.id, details=doc.doc_number)

    names = request.form.getlist('part_name[]')
    specs = request.form.getlist('part_spec[]')
    qtys  = request.form.getlist('part_qty[]')
    units = request.form.getlist('part_unit[]')
    costs = request.form.getlist('part_cost[]')

    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id = doc.id,
                name        = name.strip(),
                unit        = units[i] if i < len(units) else 'шт',
                quantity    = float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note        = specs[i] if i < len(specs) else '',
                price       = _to_float(costs[i]) if i < len(costs) else None,
            ))

    if action == 'submit' and signatory:
        db.session.add(DocumentApproval(document_id=doc.id, approver_id=signatory.id, step=0, status='pending'))

    db.session.add(DocumentComment(
        document_id = doc.id,
        author_id   = current_user.id,
        text        = f'Дефектный акт {"отправлен на согласование" if action == "submit" else "сохранён как черновик"} пользователем {current_user.full_name}.',
        is_system   = True,
    ))

    db.session.commit()

    if action == 'submit' and signatory:
        db.session.add(Notification(
            user_id = signatory.id,
            title   = f'Документ на согласование: {doc.doc_number}',
            body    = doc.title[:100],
            link    = f'/documents/{doc.id}',
            is_read = False,
        ))
        db.session.commit()

    flash(f'Документ {doc.doc_number}'
          + (f' (учётный код {doc.event_code})' if doc.event_code else '')
          + f' {"отправлен на согласование" if action == "submit" else "сохранён как черновик"}.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/trebovanie/new', methods=['GET'])
@login_required
def new_trebovanie():
    """Old entry point — everything funnels into the single full form."""
    return redirect(url_for('documents.new_trebovanie_new',
                            from_defect=request.args.get('from_defect')))


@documents.route('/trebovanie/submit', methods=['POST'])
@login_required
def submit_trebovanie():
    executor_id = request.form.get('executor_id', type=int) or current_user.id
    action = request.form.get('action', 'draft')
    related_defect_id = request.form.get('related_defect_id', type=int)
    if related_defect_id:
        rel = Document.query.filter_by(id=related_defect_id, doc_type='defect_act').first()
        related_defect_id = rel.id if rel else None
    signatory_id = request.form.get('signatory_id', type=int)

    if action == 'submit' and not signatory_id:
        action = 'draft'
        flash('Документ сохранён как черновик: не выбран Утверждающий. Укажите утверждающего и отправьте документ ещё раз.', 'warning')

    doc = Document(
        doc_type      = 'purchase_req',
        title         = request.form.get('summary', 'Требование на приобретение')[:100],
        department    = request.form.get('department', current_user.department),
        urgency       = request.form.get('urgency', 'standard'),
        purpose       = request.form.get('summary'),
        justification = request.form.get('note'),
        author_id     = current_user.id,
        executor_id   = executor_id,
        related_defect_id = related_defect_id,
        status        = 'pending' if action == 'submit' else 'draft',
    )

    needed_by = request.form.get('needed_by')
    if needed_by:
        try:
            doc.needed_by = datetime.strptime(needed_by, '%Y-%m-%d').date()
        except ValueError:
            pass

    db.session.add(doc)
    db.session.flush()
    doc.generate_number()
    log_action('document_created', 'document', doc.id, details=doc.doc_number)

    _save_form_attachments(doc)

    names = request.form.getlist('item_name[]')
    specs = request.form.getlist('item_spec[]')
    qtys  = request.form.getlist('item_qty[]')
    units = request.form.getlist('item_unit[]')
    costs = request.form.getlist('item_cost[]')

    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id = doc.id,
                name        = name.strip(),
                unit        = units[i] if i < len(units) else 'шт',
                quantity    = float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note        = specs[i] if i < len(specs) else '',
                price       = _to_float(costs[i]) if i < len(costs) else None,
            ))

    db.session.add(DocumentComment(
        document_id = doc.id,
        author_id   = current_user.id,
        text        = f'Требование {"отправлено на согласование" if action == "submit" else "сохранено как черновик"} пользователем {current_user.full_name}.',
        is_system   = True,
    ))

    _build_route(doc, action, signatory_id)

    approvers = User.query.filter(
        User.role.in_(['dept_head', 'director', 'it_admin']),
        User.is_active == True
    ).all()
    for approver in approvers:
        db.session.add(Notification(
            user_id = approver.id,
            title   = f'Новый документ на согласовании: {doc.doc_number}',
            body    = doc.title[:100],
            link    = f'/documents/{doc.id}',
            is_read = False,
        ))

    db.session.commit()
    flash(f'Документ {doc.doc_number} {"отправлен на согласование" if action == "submit" else "сохранён как черновик"}.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/trebovanie-new/new', methods=['GET'])
@login_required
def new_trebovanie_new():
    import json as _json
    from app.models import RouteTemplate
    now = datetime.now().strftime('%d.%m.%Y %H:%M')
    executors = User.query.filter_by(is_active=True).order_by(User.first_name, User.last_name).all()
    approvers = User.query.filter(User.is_active==True, User.role.in_(['dept_head', 'director', 'it_admin'])).order_by(User.first_name, User.last_name).all()
    open_defects = Document.query.filter_by(doc_type='defect_act', defect_closed=False)\
                                 .order_by(Document.created_at.desc()).limit(100).all()
    from_defect = request.args.get('from_defect', type=int)
    route_templates = RouteTemplate.query.order_by(RouteTemplate.name).all()
    templates_json = _json.dumps(
        [{'id': t.id, 'name': t.name, 'data': _json.loads(t.data)} for t in route_templates],
        ensure_ascii=False)
    return render_template('documents/trebovanie_new.html', now=now, executors=executors,
                           approvers=approvers, open_defects=open_defects,
                           from_defect=from_defect, templates_json=templates_json)


@documents.route('/trebovanie-new/submit', methods=['POST'])
@login_required
def submit_trebovanie_new():
    action = request.form.get('action', 'draft')
    executor_id = request.form.get('executor_id', type=int) or current_user.id
    signatory_id = request.form.get('signatory_id', type=int)

    if action == 'submit' and not signatory_id:
        action = 'draft'
        flash('Документ сохранён как черновик: не выбран Утверждающий. Укажите утверждающего и отправьте документ ещё раз.', 'warning')

    related_defect_id = request.form.get('related_defect_id', type=int) or None
    if related_defect_id:
        rel = Document.query.filter_by(id=related_defect_id, doc_type='defect_act').first()
        related_defect_id = rel.id if rel else None

    doc = Document(
        doc_type      = 'purchase_req',
        related_defect_id = related_defect_id,
        title         = request.form.get('summary', 'Требование на приобретение материалов')[:100],
        purpose       = request.form.get('summary'),
        justification = request.form.get('linked_to'),
        author_id     = current_user.id,
        executor_id   = executor_id,
        status        = 'pending' if action == 'submit' else 'draft',
        current_step  = 0,
    )

    needed_by = request.form.get('needed_by')
    if needed_by:
        doc.needed_by = datetime.strptime(needed_by, '%Y-%m-%d').date()

    db.session.add(doc)
    db.session.flush()
    doc.generate_number()
    log_action('document_created', 'document', doc.id, details=doc.doc_number)

    names = request.form.getlist('item_name[]')
    specs = request.form.getlist('item_spec[]')
    qtys  = request.form.getlist('item_qty[]')
    units = request.form.getlist('item_unit[]')
    dates = request.form.getlist('item_date[]')
    notes = request.form.getlist('item_note[]')
    costs = request.form.getlist('item_cost[]')

    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id = doc.id,
                name        = name.strip(),
                unit        = units[i] if i < len(units) else 'шт',
                quantity    = float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note        = notes[i] if i < len(notes) else '',
                price       = _to_float(costs[i]) if i < len(costs) else None,
            ))

    db.session.add(DocumentComment(
        document_id = doc.id,
        author_id   = current_user.id,
        text        = f'Требование {"отправлено на согласование" if action == "submit" else "сохранено как черновик"} пользователем {current_user.full_name}.',
        is_system   = True,
    ))

    _build_route(doc, action, signatory_id)
    _save_form_attachments(doc)

    db.session.commit()

    if action == 'submit':
        _notify_approvers(doc)

    db.session.commit()
    flash(f'Документ {doc.doc_number} {"отправлен на согласование" if action == "submit" else "сохранён как черновик"}.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/po-services/new', methods=['GET'])
@login_required
def new_po_services():
    from_req = request.args.get('from_req', type=int)
    linked_req = None
    if from_req:
        linked_req = Document.query.filter_by(id=from_req, doc_type='purchase_req').first()
    executors = User.query.filter_by(is_active=True).order_by(User.first_name, User.last_name).all()
    return render_template('documents/po_services.html', executors=executors,
                           linked_req=linked_req)


@documents.route('/po-services/submit', methods=['POST'])
@login_required
def submit_po_services():
    action = request.form.get('action', 'draft')
    signatory_id = request.form.get('signatory_id', type=int)

    if action == 'submit' and not signatory_id:
        action = 'draft'
        flash('Документ сохранён как черновик: не выбран Утверждающий. Укажите утверждающего и отправьте документ ещё раз.', 'warning')

    extras = {
        'Заказано': request.form.get('ordered_by'),
        'Плательщик НДС': request.form.get('vat_payer'),
        'Валюта': request.form.get('currency'),
        'Дата оплаты': request.form.get('payment_date'),
        'Основание': request.form.get('basis'),
        'Контрагент': request.form.get('counterparty'),
        'Условия оплаты': request.form.get('payment_terms'),
        'Условия оказания услуг': request.form.get('service_terms'),
        'Срок оказания услуг': request.form.get('service_date'),
        'Статья бюджета': request.form.get('budget_line'),
    }
    extras_text = '\n'.join(f'{k}: {v}' for k, v in extras.items() if v)

    doc = Document(
        doc_type      = 'po_services',
        related_req_id = (request.form.get('related_req_id', type=int) or None),
        title         = request.form.get('summary', 'РО на услуги')[:100],
        department    = request.form.get('department', current_user.department),
        purpose       = request.form.get('summary'),
        justification = extras_text,
        author_id     = current_user.id,
        status        = 'pending' if action == 'submit' else 'draft',
        current_step  = 0,
    )

    service_date = request.form.get('service_date')
    if service_date:
        try:
            doc.needed_by = datetime.strptime(service_date, '%Y-%m-%d').date()
        except ValueError:
            pass

    db.session.add(doc)
    db.session.flush()
    doc.generate_number()
    log_action('document_created', 'document', doc.id, details=doc.doc_number)

    names = request.form.getlist('item_name[]')
    units = request.form.getlist('item_unit[]')
    qtys  = request.form.getlist('item_qty[]')
    costs = request.form.getlist('item_cost[]')

    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id = doc.id,
                name        = name.strip(),
                unit        = units[i] if i < len(units) else '',
                quantity    = float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note        = '',
                price       = _to_float(costs[i]) if i < len(costs) else None,
            ))

    db.session.add(DocumentComment(
        document_id = doc.id,
        author_id   = current_user.id,
        text        = f'РО на услуги {"отправлен на согласование" if action == "submit" else "сохранён как черновик"} пользователем {current_user.full_name}.',
        is_system   = True,
    ))

    _build_route(doc, action, signatory_id)
    _save_form_attachments(doc)

    req_id = request.form.get('related_req_id', type=int)
    if req_id:
        req = Document.query.filter_by(id=req_id, doc_type='purchase_req').first()
        if req and req.status in ('approved', 'in_execution'):
            req.status = 'in_execution'
            db.session.add(DocumentComment(
                document_id=req.id, author_id=current_user.id,
                text=f'Создан ПО по требованию (см. связанные документы).',
                is_system=True))
    db.session.commit()

    if action == 'submit':
        _notify_approvers(doc)

    db.session.commit()
    flash(f'Документ {doc.doc_number} {"отправлен на согласование" if action == "submit" else "сохранён как черновик"}.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/po-services/<int:doc_id>/edit', methods=['GET'])
@login_required
def edit_po_services(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'po_services' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    extras = _unpack_extras(doc.justification)
    items = doc.items.all()
    return render_template('documents/po_services_edit.html', doc=doc, extras=extras, items=items)


@documents.route('/po-services/<int:doc_id>/update', methods=['POST'])
@login_required
def update_po_services(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'po_services' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))

    doc.title      = request.form.get('summary', doc.title)[:100]
    doc.purpose    = request.form.get('summary')
    doc.department = request.form.get('department', doc.department)

    extras = {
        'Заказано': request.form.get('ordered_by'),
        'Плательщик НДС': request.form.get('vat_payer'),
        'Валюта': request.form.get('currency'),
        'Дата оплаты': request.form.get('payment_date'),
        'Основание': request.form.get('basis'),
        'Контрагент': request.form.get('counterparty'),
        'Условия оплаты': request.form.get('payment_terms'),
        'Условия оказания услуг': request.form.get('service_terms'),
        'Срок оказания услуг': request.form.get('service_date'),
        'Статья бюджета': request.form.get('budget_line'),
    }
    doc.justification = '\n'.join(f'{k}: {v}' for k, v in extras.items() if v)

    service_date = request.form.get('service_date')
    if service_date:
        try:
            doc.needed_by = datetime.strptime(service_date, '%Y-%m-%d').date()
        except ValueError:
            pass

    for it in doc.items.all():
        db.session.delete(it)
    names = request.form.getlist('item_name[]')
    units = request.form.getlist('item_unit[]')
    qtys  = request.form.getlist('item_qty[]')
    costs = request.form.getlist('item_cost[]')
    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id=doc.id,
                name=name.strip(),
                unit=units[i] if i < len(units) else '',
                quantity=float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note='',
                price=_to_float(costs[i]) if i < len(costs) else None,
            ))

    _save_form_attachments(doc)

    action = request.form.get('action', 'save')
    if action == 'submit':
        doc.status = 'pending'
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ отредактирован и повторно отправлен на согласование пользователем {current_user.full_name}.',
            is_system=True,
        ))
        _notify_current_approvers(doc, f'Документ повторно на согласовании: {doc.doc_number}')
        msg = 'Документ отредактирован и отправлен на согласование.'
    else:
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ обновлён пользователем {current_user.full_name} (черновик доработки).',
            is_system=True,
        ))
        msg = 'Изменения сохранены. Документ всё ещё на доработке.'

    db.session.commit()
    flash(msg, 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/po-trebovanie/new', methods=['GET'])
@login_required
def new_po_trebovanie():
    executors = User.query.filter_by(is_active=True).order_by(User.first_name, User.last_name).all()
    return render_template('documents/po_trebovanie.html', executors=executors)


@documents.route('/po-trebovanie/submit', methods=['POST'])
@login_required
def submit_po_trebovanie():
    action = request.form.get('action', 'draft')
    signatory_id = request.form.get('signatory_id', type=int)

    if action == 'submit' and not signatory_id:
        action = 'draft'
        flash('Документ сохранён как черновик: не выбран Утверждающий. Укажите утверждающего и отправьте документ ещё раз.', 'warning')

    extras = {
        'Основание': request.form.get('basis'),
        'Требование на приобретение материалов': request.form.get('material_request'),
        'Контрагент': request.form.get('counterparty'),
        'Плательщик НДС': request.form.get('vat_payer'),
        'Заказано': request.form.get('ordered_by'),
        'Валюта': request.form.get('currency'),
        'Дата оплаты': request.form.get('payment_date'),
        'Условия оплаты': request.form.get('payment_terms'),
        'Условия оказания услуг': request.form.get('service_terms'),
        'Срок оказания услуг': request.form.get('service_date'),
        'Статья бюджета': request.form.get('budget_line'),
    }
    extras_text = '\n'.join(f'{k}: {v}' for k, v in extras.items() if v)

    doc = Document(
        doc_type      = 'po_trebovanie',
        title         = request.form.get('summary', 'РО на товары')[:100],
        department    = request.form.get('department', current_user.department),
        purpose       = request.form.get('summary'),
        justification = extras_text,
        author_id     = current_user.id,
        status        = 'pending' if action == 'submit' else 'draft',
        current_step  = 0,
    )

    payment_date = request.form.get('payment_date')
    if payment_date:
        try:
            doc.needed_by = datetime.strptime(payment_date, '%Y-%m-%d').date()
        except ValueError:
            pass

    db.session.add(doc)
    db.session.flush()
    doc.generate_number()
    log_action('document_created', 'document', doc.id, details=doc.doc_number)

    names = request.form.getlist('item_name[]')
    units = request.form.getlist('item_unit[]')
    qtys  = request.form.getlist('item_qty[]')
    costs = request.form.getlist('item_cost[]')
    codes = request.form.getlist('item_code[]')

    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id = doc.id,
                name        = name.strip(),
                unit        = units[i] if i < len(units) else '',
                quantity    = float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note        = codes[i] if i < len(codes) else '',
                price       = _to_float(costs[i]) if i < len(costs) else None,
            ))

    db.session.add(DocumentComment(
        document_id = doc.id,
        author_id   = current_user.id,
        text        = f'РО на товары {"отправлена на согласование" if action == "submit" else "сохранена как черновик"} пользователем {current_user.full_name}.',
        is_system   = True,
    ))

    _build_route(doc, action, signatory_id)
    _save_form_attachments(doc)

    db.session.commit()

    if action == 'submit':
        _notify_approvers(doc)

    db.session.commit()
    flash(f'Документ {doc.doc_number} {"отправлен на согласование" if action == "submit" else "сохранён как черновик"}.', 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/po-trebovanie/<int:doc_id>/edit', methods=['GET'])
@login_required
def edit_po_trebovanie(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'po_trebovanie' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    extras = _unpack_extras(doc.justification)
    items = doc.items.all()
    return render_template('documents/po_trebovanie_edit.html', doc=doc, extras=extras, items=items)


@documents.route('/po-trebovanie/<int:doc_id>/update', methods=['POST'])
@login_required
def update_po_trebovanie(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'po_trebovanie' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))

    doc.title      = request.form.get('summary', doc.title)[:100]
    doc.purpose    = request.form.get('summary')
    doc.department = request.form.get('department', doc.department)

    extras = {
        'Основание': request.form.get('basis'),
        'Требование на приобретение материалов': request.form.get('material_request'),
        'Контрагент': request.form.get('counterparty'),
        'Плательщик НДС': request.form.get('vat_payer'),
        'Заказано': request.form.get('ordered_by'),
        'Валюта': request.form.get('currency'),
        'Дата оплаты': request.form.get('payment_date'),
        'Условия оплаты': request.form.get('payment_terms'),
        'Условия оказания услуг': request.form.get('service_terms'),
        'Срок оказания услуг': request.form.get('service_date'),
        'Статья бюджета': request.form.get('budget_line'),
    }
    doc.justification = '\n'.join(f'{k}: {v}' for k, v in extras.items() if v)

    payment_date = request.form.get('payment_date')
    if payment_date:
        try:
            doc.needed_by = datetime.strptime(payment_date, '%Y-%m-%d').date()
        except ValueError:
            pass

    for it in doc.items.all():
        db.session.delete(it)
    names = request.form.getlist('item_name[]')
    units = request.form.getlist('item_unit[]')
    qtys  = request.form.getlist('item_qty[]')
    costs = request.form.getlist('item_cost[]')
    codes = request.form.getlist('item_code[]')
    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id=doc.id,
                name=name.strip(),
                unit=units[i] if i < len(units) else '',
                quantity=float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note=codes[i] if i < len(codes) else '',
                price=_to_float(costs[i]) if i < len(costs) else None,
            ))

    _save_form_attachments(doc)

    action = request.form.get('action', 'save')
    if action == 'submit':
        doc.status = 'pending'
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ отредактирован и повторно отправлен на согласование пользователем {current_user.full_name}.',
            is_system=True,
        ))
        _notify_current_approvers(doc, f'Документ повторно на согласовании: {doc.doc_number}')
        msg = 'Документ отредактирован и отправлен на согласование.'
    else:
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ обновлён пользователем {current_user.full_name} (черновик доработки).',
            is_system=True,
        ))
        msg = 'Изменения сохранены. Документ всё ещё на доработке.'

    db.session.commit()
    flash(msg, 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/trebovanie/<int:doc_id>/edit', methods=['GET'])
@login_required
def edit_trebovanie(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'trebovanie' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    items = doc.items.all()
    return render_template('documents/trebovanie_edit.html', doc=doc, items=items)


@documents.route('/trebovanie/<int:doc_id>/update', methods=['POST'])
@login_required
def update_trebovanie(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'trebovanie' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))

    doc.title   = request.form.get('summary', doc.title)[:100]
    doc.purpose = request.form.get('summary')
    doc.justification = request.form.get('linked_to')

    needed_by = request.form.get('needed_by')
    if needed_by:
        try:
            doc.needed_by = datetime.strptime(needed_by, '%Y-%m-%d').date()
        except ValueError:
            pass

    for it in doc.items.all():
        db.session.delete(it)
    names = request.form.getlist('item_name[]')
    units = request.form.getlist('item_unit[]')
    qtys  = request.form.getlist('item_qty[]')
    notes = request.form.getlist('item_note[]')
    costs = request.form.getlist('item_cost[]')
    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id=doc.id,
                name=name.strip(),
                unit=units[i] if i < len(units) else 'шт',
                quantity=float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note=notes[i] if i < len(notes) else '',
                price=_to_float(costs[i]) if i < len(costs) else None,
            ))

    _save_form_attachments(doc)

    action = request.form.get('action', 'save')
    if action == 'submit':
        doc.status = 'pending'
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ отредактирован и повторно отправлен на согласование пользователем {current_user.full_name}.',
            is_system=True,
        ))
        _notify_current_approvers(doc, f'Документ повторно на согласовании: {doc.doc_number}')
        msg = 'Документ отредактирован и отправлен на согласование.'
    else:
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ обновлён пользователем {current_user.full_name} (черновик доработки).',
            is_system=True,
        ))
        msg = 'Изменения сохранены. Документ всё ещё на доработке.'

    db.session.commit()
    flash(msg, 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/defect-act/<int:doc_id>/edit', methods=['GET'])
@login_required
def edit_defect_act(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'defect_act' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    items = doc.items.all()
    return render_template('documents/defect_edit.html', doc=doc, items=items)


@documents.route('/defect-act/<int:doc_id>/update', methods=['POST'])
@login_required
def update_defect_act(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.author_id != current_user.id:
        abort(403)
    if doc.doc_type != 'defect_act' or doc.status != 'returned':
        flash('Этот документ нельзя редактировать.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))

    doc.title         = request.form.get('description', doc.title)[:100]
    doc.purpose       = request.form.get('description')
    doc.justification = request.form.get('cause')
    doc.urgency       = request.form.get('urgency', doc.urgency)

    for it in doc.items.all():
        db.session.delete(it)
    names = request.form.getlist('part_name[]')
    specs = request.form.getlist('part_spec[]')
    qtys  = request.form.getlist('part_qty[]')
    units = request.form.getlist('part_unit[]')
    costs = request.form.getlist('part_cost[]')
    for i, name in enumerate(names):
        if name.strip():
            db.session.add(DocumentItem(
                document_id=doc.id,
                name=name.strip(),
                unit=units[i] if i < len(units) else 'шт',
                quantity=float(qtys[i]) if i < len(qtys) and qtys[i] else None,
                note=specs[i] if i < len(specs) else '',
                price=_to_float(costs[i]) if i < len(costs) else None,
            ))

    _save_form_attachments(doc)

    action = request.form.get('action', 'save')
    if action == 'submit':
        doc.status = 'pending'
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ отредактирован и повторно отправлен на согласование пользователем {current_user.full_name}.',
            is_system=True,
        ))
        _notify_current_approvers(doc, f'Документ повторно на согласовании: {doc.doc_number}')
        msg = 'Документ отредактирован и отправлен на согласование.'
    else:
        db.session.add(DocumentComment(
            document_id=doc.id, author_id=current_user.id,
            text=f'Документ обновлён пользователем {current_user.full_name} (черновик доработки).',
            is_system=True,
        ))
        msg = 'Изменения сохранены. Документ всё ещё на доработке.'

    db.session.commit()
    flash(msg, 'success')
    return redirect(url_for('documents.view', doc_id=doc.id))


@documents.route('/<int:doc_id>/print')
@login_required
def print_doc(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if (not current_user.can_access('documents')
            and doc.author_id != current_user.id
            and not _is_assigned_approver(doc, current_user)):
        abort(403)
    items = doc.items.all()
    extras = _unpack_extras(doc.justification)
    approvals = doc.approvals.order_by(DocumentApproval.step).all()
    return render_template('documents/print.html', doc=doc, items=items,
                           extras=extras, approvals=approvals)

@documents.route('/<int:doc_id>/close-defect', methods=['POST'])
@login_required
def close_defect(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.doc_type != 'defect_act' or doc.defect_closed:
        flash('Этот документ нельзя закрыть.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    if not (doc.author_id == current_user.id
            or current_user.can_access('documents')
            or current_user.can_access('equipment')):
        abort(403)
    doc.defect_closed = True
    doc.defect_closed_at = datetime.utcnow()
    db.session.add(DocumentComment(
        document_id=doc.id, author_id=current_user.id,
        text=f'Дефект {doc.event_code or doc.doc_number} закрыт (ремонт выполнен).',
        is_system=True))
    log_action('defect_closed', 'document', doc.id, details=doc.event_code or doc.doc_number)
    db.session.commit()
    flash(f'Дефект {doc.event_code or doc.doc_number} закрыт.', 'success')
    return redirect(url_for('documents.view', doc_id=doc_id))


@documents.route('/<int:doc_id>/mark-executed', methods=['POST'])
@login_required
def mark_executed(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.doc_type != 'purchase_req' or doc.status not in ('approved', 'in_execution'):
        flash('Документ нельзя отметить исполненным.', 'warning')
        return redirect(url_for('documents.view', doc_id=doc_id))
    if not (current_user.role in ('procurement', 'it_admin', 'director')
            or current_user.can_access('documents_procurement')):
        abort(403)
    doc.status = 'executed'
    db.session.add(DocumentComment(
        document_id=doc.id, author_id=current_user.id,
        text=f'Требование исполнено ({current_user.full_name}).', is_system=True))
    db.session.add(Notification(
        user_id=doc.author_id, title=f'Требование {doc.doc_number} исполнено',
        body=(doc.title or '')[:100], link=f'/documents/{doc.id}', is_read=False))
    log_action('document_executed', 'document', doc.id, details=doc.doc_number)
    db.session.commit()
    flash('Требование отмечено исполненным.', 'success')
    return redirect(url_for('documents.view', doc_id=doc_id))


@documents.route('/route-templates/save', methods=['POST'])
@login_required
def save_route_template():
    import json as _json
    from app.models import RouteTemplate
    name = (request.form.get('template_name') or '').strip()
    raw = request.form.get('template_json') or ''
    if not name or not raw:
        return jsonify({'error': 'name and route required'}), 400
    try:
        data = _json.loads(raw)
        assert isinstance(data.get('stages'), list)
    except Exception:
        return jsonify({'error': 'bad route data'}), 400
    tpl = RouteTemplate.query.filter_by(name=name).first()
    if tpl is None:
        tpl = RouteTemplate(name=name, created_by=current_user.id)
        db.session.add(tpl)
    tpl.data = _json.dumps(data, ensure_ascii=False)
    log_action('route_template_saved', details=name)
    db.session.commit()
    return jsonify({'ok': True, 'id': tpl.id, 'name': tpl.name})


@documents.route('/export.xlsx')
@login_required
def export_xlsx():
    """Flat Excel export: one row per item, with the document number.
    Filters: ?doc_type=&year=2026&month=7 (all optional; no type = all types)."""
    if not (current_user.can_access('documents') or current_user.can_access('documents_read')
            or current_user.can_access('documents_procurement')):
        abort(403)
    import io
    from openpyxl import Workbook
    from sqlalchemy import extract

    year = request.args.get('year', type=int)
    month = request.args.get('month', type=int)
    doc_type = request.args.get('doc_type') or None
    if doc_type and doc_type not in DOC_TYPES:
        abort(400)

    query = Document.query
    if doc_type:
        query = query.filter_by(doc_type=doc_type)
    if year:
        query = query.filter(extract('year', Document.created_at) == year)
    if month:
        query = query.filter(extract('month', Document.created_at) == month)
    docs = query.order_by(Document.created_at).all()

    wb = Workbook()
    ws = wb.active
    ws.title = (DOC_TYPES.get(doc_type, 'Документы')[:31] if doc_type else 'Документы')
    ws.append(['Номер', 'Тип', 'Дата', 'Статус', 'Отдел', 'Автор', 'Код ДА',
               'Наименование', 'Характеристика/прим.', 'Кол-во', 'Ед.',
               'Цена', 'Сумма'])
    from app.models import DOC_STATUSES
    for doc in docs:
        items = doc.items.all() or [None]
        for it in items:
            ws.append([
                doc.doc_number,
                DOC_TYPES.get(doc.doc_type, doc.doc_type),
                doc.created_at.strftime('%d.%m.%Y'),
                DOC_STATUSES.get(doc.status, doc.status),
                doc.department or '',
                doc.author.full_name if doc.author else '',
                (doc.related_defect.event_code if doc.related_defect else '') or '',
                it.name if it else '',
                (it.note or '') if it else '',
                float(it.quantity) if it and it.quantity is not None else None,
                (it.unit or '') if it else '',
                float(it.price) if it and it.price is not None else None,
                (float(it.quantity) * float(it.price))
                    if it and it.quantity is not None and it.price is not None else None,
            ])
    for col, width in zip('ABCDEFGHIJKLM', [16, 22, 11, 14, 12, 20, 12, 40, 24, 8, 6, 12, 14]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    suffix = f'-{doc_type}' if doc_type else ''
    suffix += f'-{year}' if year else ''
    suffix += f'-{month:02d}' if month else ''
    from flask import send_file
    log_action('documents_exported',
               details=f'type={doc_type or "all"} year={year} month={month} docs={len(docs)}')
    db.session.commit()
    return send_file(buf, as_attachment=True,
                     download_name=f'documents{suffix}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
